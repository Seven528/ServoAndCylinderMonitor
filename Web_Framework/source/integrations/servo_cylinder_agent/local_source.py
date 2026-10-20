"""Local CSV/XLSX streaming. File size and record count are not capped."""
from __future__ import annotations
import codecs
from itertools import islice
from pathlib import Path
from zipfile import BadZipFile
import pandas as pd
from openpyxl import load_workbook

BATCH_ROWS = 10_000


def validate_file(path):
    path = Path(path).expanduser()
    if not path.is_absolute():
        raise ValueError('请输入完整的绝对文件路径，或点击选择文件。')
    path = path.resolve()
    if path.suffix.lower() not in ('.csv', '.xlsx'):
        raise ValueError('本地分析支持 CSV 或 XLSX 文件。')
    if not path.is_file():
        raise ValueError('该文件不存在或不是普通文件，请检查文件路径。')
    return path


def inspect_file(path):
    path = validate_file(path)
    sheets = []
    if path.suffix.lower() == '.xlsx':
        try:
            book = load_workbook(path,read_only=True,data_only=True,keep_links=False)
            try:sheets=book.sheetnames
            finally:book.close()
        except (BadZipFile,KeyError,OSError):
            raise ValueError('无法打开 XLSX，请确认文件格式正确、没有加密且未损坏。') from None
        if not sheets:
            raise ValueError('文件没有工作表。')
    return {'name':path.name,'bytes':path.stat().st_size,'sheets':sheets}


def _csv_encoding(path):
    # Validate UTF-8 incrementally before ingestion, so a late decoding error
    # cannot restart the source and count an already processed prefix twice.
    decoder=codecs.getincrementaldecoder('utf-8-sig')()
    try:
        with path.open('rb') as stream:
            while True:
                data=stream.read(1024*1024)
                if not data:break
                decoder.decode(data,final=False)
            decoder.decode(b'',final=True)
        return 'utf-8-sig'
    except UnicodeDecodeError:
        return 'gb18030'


def frame_stream(path,sheet=None,batch_rows=BATCH_ROWS):
    path=validate_file(path)
    if not isinstance(batch_rows,int) or batch_rows<1:
        raise ValueError('批次行数须为正整数。')
    if path.suffix.lower()=='.csv':
        if sheet is not None:raise ValueError('CSV 文件不支持工作表选择。')
        try:
            with pd.read_csv(path,encoding=_csv_encoding(path),dtype=object,chunksize=batch_rows) as reader:
                for frame in reader:
                    if not frame.empty:yield frame
        except UnicodeDecodeError:
            raise ValueError('无法识别 CSV 编码，请另存为 UTF-8 CSV。') from None
        return
    book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        selected=sheet if sheet is not None else book.sheetnames[0]
        if selected not in book.sheetnames:raise ValueError('所选工作表不存在，请重新选择文件。')
        rows=book[selected].iter_rows(values_only=True)
        header=next(rows,None)
        if not header:raise ValueError('工作表没有字段标题。')
        names=[str(value).strip() if value is not None else f'Unnamed: {i}' for i,value in enumerate(header)]
        while True:
            batch=list(islice(rows,batch_rows))
            if not batch:break
            batch=[row for row in batch if any(value is not None for value in row)]
            if batch:yield pd.DataFrame(batch,columns=names)
    finally:book.close()


def local_stream(path,sheet=None):
    path=validate_file(path)
    details=inspect_file(path)
    selected=sheet if sheet is not None else (details['sheets'][0] if details['sheets'] else None)
    source={'kind':'xlsx' if details['sheets'] else 'csv','filename':path.name,'sheet':selected,'unlimited':True}
    return frame_stream(path,selected),source


def load_frame(path,sheet=None):
    """Compatibility helper; production analysis consumes frame_stream instead."""
    chunks,source=local_stream(path,sheet)
    frames=list(chunks)
    if not frames:raise ValueError('文件或工作表没有数据。')
    return pd.concat(frames,ignore_index=True),source
